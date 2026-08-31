from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
OUT_DIR.mkdir(exist_ok=True)
OUT_PATH = OUT_DIR / "Hotel_ElbRivera_Digital_Growth_Proposal_Veer_Inderpreet.docx"

# Resolved design preset: narrative_proposal.
# Named brand overrides: River Navy, Elbe Blue, Warm Gold, Sand, Mist.
FONT = "Calibri"
NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GOLD = "B98A3D"
SAND = "F4EEE4"
MIST = "F4F6F9"
PALE_BLUE = "E8EEF5"
INK = "203040"
MUTED = "65717D"
WHITE = "FFFFFF"
GREEN = "3E6B58"
RED = "9B1C1C"
BORDER = "D5DEE7"

PAGE_W_DXA = 12240
PAGE_H_DXA = 15840
CONTENT_W_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=BORDER, size=6, inside=True):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    names = ["top", "left", "bottom", "right"]
    if inside:
        names += ["insideH", "insideV"]
    for name in names:
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    if sum(widths_dxa) != CONTENT_W_DXA:
        raise ValueError(f"Table columns must total {CONTENT_W_DXA} DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_W_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run(run, *, size=11, color=INK, bold=False, italic=False, font=FONT):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic
    return run


def keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_keep_lines(paragraph, value=True):
    paragraph.paragraph_format.keep_together = value


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True, size=10):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_node = OxmlElement("w:t")
    text_node.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text_node, end])
    set_run(run, size=8.5, color=MUTED)


def add_page_number(paragraph):
    paragraph.add_run("Page ")
    add_field(paragraph, "PAGE")
    paragraph.add_run(" of ")
    add_field(paragraph, "NUMPAGES")
    for run in paragraph.runs:
        set_run(run, size=8.5, color=MUTED)


def create_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs, default=-1) + 1
    next_num = max(existing_num, default=0) + 1

    def define(fmt, text, font=None):
        nonlocal next_abs, next_num
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(next_abs))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "280")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "290")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        lvl.extend([start, num_fmt, lvl_text, suff, lvl_jc, p_pr])
        if font:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), font)
            r_fonts.set(qn("w:hAnsi"), font)
            r_pr.append(r_fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        first_num = numbering.find(qn("w:num"))
        if first_num is None:
            numbering.append(abstract)
        else:
            numbering.insert(list(numbering).index(first_num), abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(next_num))
        abstract_id = OxmlElement("w:abstractNumId")
        abstract_id.set(qn("w:val"), str(next_abs))
        num.append(abstract_id)
        numbering.append(num)
        result = next_num
        next_abs += 1
        next_num += 1
        return result

    return define("bullet", "•", "Arial"), define("decimal", "%1.")


def set_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def add_bullet(doc, text, bullet_id, *, bold_lead=None, color=INK, size=10.5, after=4):
    p = doc.add_paragraph()
    set_num(p, bullet_id)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.208
    p.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), size=size, color=color, bold=True)
        set_run(p.add_run(text[len(bold_lead):]), size=size, color=color)
    else:
        set_run(p.add_run(text), size=size, color=color)
    return p


def add_numbered(doc, text, number_id, *, bold_lead=None):
    p = doc.add_paragraph()
    set_num(p, number_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    p.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), size=10.5, bold=True)
        set_run(p.add_run(text[len(bold_lead):]), size=10.5)
    else:
        set_run(p.add_run(text), size=10.5)
    return p


def add_kicker(doc, text, color=GOLD, align=WD_ALIGN_PARAGRAPH.LEFT, after=5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    set_run(r, size=8.5, color=color, bold=True)
    r.font.letter_spacing = Pt(1.1) if hasattr(r.font, "letter_spacing") else None
    return p


def add_title(doc, text, size=27, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT, after=8):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), size=size, color=color, bold=True)
    return p


def add_subtitle(doc, text, size=13, color=MUTED, align=WD_ALIGN_PARAGRAPH.LEFT, after=14):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), size=size, color=color)
    return p


def add_body(doc, text, *, bold_lead=None, after=8, size=10.6, align=WD_ALIGN_PARAGRAPH.JUSTIFY, color=INK):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.28
    set_keep_lines(p)
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), size=size, color=color, bold=True)
        set_run(p.add_run(text[len(bold_lead):]), size=size, color=color)
    else:
        set_run(p.add_run(text), size=size, color=color)
    return p


def add_section_title(doc, kicker, title, subtitle=None):
    add_kicker(doc, kicker)
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = False
    p.paragraph_format.keep_with_next = True
    p.add_run(title)
    if subtitle:
        add_subtitle(doc, subtitle, size=11.3, after=12)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_h3(doc, text):
    p = doc.add_paragraph(style="Heading 3")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_callout(doc, label, text, *, fill=SAND, accent=GOLD):
    table = doc.add_table(rows=1, cols=1)
    # Named callout override: 190 DXA inset and matching start cell margin.
    set_table_geometry(table, [CONTENT_W_DXA], indent_dxa=190)
    set_table_borders(table, color=fill, size=0, inside=False)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=190, bottom=150, end=190)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(label.upper() + "  "), size=8.8, color=accent, bold=True)
    set_run(p.add_run(text), size=11.2, color=NAVY, bold=True)
    return table


def add_spacer(doc, points=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1
    r = p.add_run(" ")
    r.font.size = Pt(points)
    return p


def add_page_break(doc):
    doc.add_page_break()


def add_matrix(doc, headers, rows, widths, *, header_fill=NAVY, body_fill=WHITE, font_size=9.4):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table, color=BORDER, size=6, inside=True)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, label in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(label), size=9.2, color=WHITE, bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            cell = cells[idx]
            set_cell_shading(cell, body_fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            set_run(p.add_run(str(value)), size=font_size, color=INK)
    set_table_geometry(table, widths)
    return table


def set_doc_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    style_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in style_tokens.items():
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name in ("Pitch Small", "Pitch Table"):
        if style_name not in [s.name for s in doc.styles]:
            style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(9.2)
        style.font.color.rgb = rgb(INK)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15

    core = doc.core_properties
    core.title = "Hotel ElbRivera Digital Growth Proposal"
    core.subject = "Website, SEO, GEO, Google visibility, reviews, ads, content and social media"
    core.author = "Veer Pratap Singh and Inderpreet Singh"
    core.keywords = "Hotel ElbRivera, hospitality marketing, SEO, GEO, Google Ads, Google Business Profile"
    core.comments = "Prepared as a client proposal based on a public website audit."


def set_headers_footers(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    first_header = section.first_page_header
    p = first_header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("CONFIDENTIAL PROPOSAL"), size=7.5, color=MUTED, bold=True)

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [4680, 4680], indent_dxa=0)
    set_table_borders(table, color=WHITE, size=0, inside=False)
    p1 = table.cell(0, 0).paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    set_run(p1.add_run("V+I  |  HOTEL ELBRIVERA"), size=8, color=NAVY, bold=True)
    p2 = table.cell(0, 1).paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.space_after = Pt(0)
    set_run(p2.add_run("DIGITAL GROWTH PROPOSAL"), size=8, color=MUTED, bold=True)

    first_footer = section.first_page_footer
    fp = first_footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(0)
    set_run(fp.add_run("Veer Pratap Singh + Inderpreet Singh  |  31 August 2026"), size=8.2, color=MUTED)

    footer = section.footer
    ft = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(ft, [6500, 2860], indent_dxa=0)
    set_table_borders(ft, color=WHITE, size=0, inside=False)
    left = ft.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    set_run(left.add_run("Web, content, brand and growth partners"), size=8.5, color=MUTED)
    right = ft.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    add_page_number(right)


def build_document():
    doc = Document()
    set_doc_styles(doc)
    set_headers_footers(doc)
    bullet_id, number_id = create_numbering(doc)

    # COVER
    add_spacer(doc, 42)
    add_kicker(doc, "Digital growth proposal", color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    add_title(doc, "Hotel ElbRivera", size=34, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    add_subtitle(doc, "From a useful website to a connected direct-booking and demand-generation system", size=15, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    add_callout(doc, "The opportunity", "Turn the hotel's riverside location, restaurant, events and meeting facilities into more direct discovery, enquiries and bookings.", fill=SAND, accent=GOLD)
    add_spacer(doc, 18)
    cover_meta = doc.add_table(rows=3, cols=2)
    set_table_geometry(cover_meta, [4680, 4680])
    set_table_borders(cover_meta, color=WHITE, size=0, inside=False)
    meta = [
        ("PREPARED FOR", "Hotel ElbRivera, Magdeburg"),
        ("PREPARED BY", "Veer Pratap Singh + Inderpreet Singh"),
        ("SCOPE", "Website • Search • Maps • Reviews • GEO • Ads • Content"),
    ]
    for row, (label, value) in zip(cover_meta.rows, meta):
        set_cell_shading(row.cells[0], PALE_BLUE)
        set_cell_shading(row.cells[1], MIST)
        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        set_run(p0.add_run(label), size=8.5, color=BLUE, bold=True)
        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_run(p1.add_run(value), size=10.2, color=NAVY, bold=True)
    add_spacer(doc, 22)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run("V+I"), size=18, color=GOLD, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    add_hyperlink(p, "veer-pratap-singh-portfolio.vercel.app", "https://veer-pratap-singh-portfolio.vercel.app", color=BLUE)

    # PAGE 2
    add_page_break(doc)
    add_section_title(doc, "Executive proposal", "The property is stronger than its digital presentation", "ElbRivera does not lack services or content. It lacks prioritisation, contemporary presentation and a measurable conversion system.")
    add_callout(doc, "Core recommendation", "Build one guest journey around Stay, Eat, Celebrate and Meet - then support it with local visibility, reviews, SEO, GEO, paid media and monthly content.", fill=PALE_BLUE, accent=BLUE)
    add_h2(doc, "What we would help ElbRivera achieve")
    outcomes = [
        "Make the hotel's calm riverside setting and proximity to Magdeburg immediately understandable.",
        "Increase qualified direct-booking clicks and reduce dependence on fragmented guest journeys.",
        "Turn restaurant events, celebrations and meetings into clearly packaged, trackable offers.",
        "Strengthen visibility across Google Search, Google Maps, Google hotel results and AI-assisted discovery.",
        "Create a consistent review, content and advertising operation that remains active after launch.",
    ]
    for item in outcomes:
        add_bullet(doc, item, bullet_id)
    add_h2(doc, "The promise we will make - and the one we will not")
    add_body(doc, "We will build, manage and measure the factors within our control: website quality, content usefulness, local profile completeness, review workflows, technical SEO, campaign relevance and conversion tracking. No responsible provider can guarantee a permanent #1 organic position, a particular star rating or inclusion in an AI answer. Our commitment is disciplined improvement, transparent reporting and better commercial readiness.")

    # PAGE 3
    add_page_break(doc)
    add_section_title(doc, "Current-state audit", "Where the existing website is losing momentum", "Public website review completed 30 August 2026. Internal booking, revenue and analytics data were not available, so commercial impact is directional until baseline measurement is installed.")
    audit_rows = [
        ("Events", "Old dates, a visible raw search phrase, small images and no clear event-level reservation journey.", "Reduced trust and missed restaurant demand."),
        ("Enquiries", "One broad form contains roughly 74 controls across stays, restaurant, wellness, meetings and equipment.", "High cognitive load and likely abandonment."),
        ("Navigation", "A deep hierarchy spreads rooms, packages, restaurant, events, meetings and destination content across many routes.", "Guests must work to find the next action."),
        ("Booking", "The external booking engine is functional but visually disconnected; dates are not strongly selected before handoff.", "Lost context at the moment of highest intent."),
        ("Content", "Repetition, spelling issues, stale information and inconsistent claims appear across public pages.", "Lower confidence for guests and organisers."),
        ("Mobile", "PageSpeed showed 79 performance and 89 accessibility for the audited event page, while desktop performance was 96.", "A reasonable technical base, but a weaker mobile sales experience."),
    ]
    add_matrix(doc, ["AREA", "OBSERVATION", "COMMERCIAL EFFECT"], audit_rows, [1500, 4660, 3200], font_size=8.8)
    add_spacer(doc, 8)
    add_h2(doc, "What should be preserved")
    for item in [
        "A distinctive quiet setting by the Elbe and the Elberadweg.",
        "Direct-booking savings, free parking, Wi-Fi and charging facilities.",
        "Multiple revenue lines: rooms, restaurant, wellness, events, celebrations and meetings.",
        "Existing booking technology and generally sound desktop performance and basic SEO setup.",
    ]:
        add_bullet(doc, item, bullet_id)

    # PAGE 4
    add_page_break(doc)
    add_section_title(doc, "Website and conversion", "A new digital guest journey", "The redesign should make the correct next step obvious within seconds, on every screen.")
    journey_rows = [
        ("STAY", "Rooms, packages, direct-booking benefit", "Check dates"),
        ("EAT", "Restaurant, menu, hours, seasonal dining", "Reserve a table"),
        ("CELEBRATE", "Weddings, family events, seasonal occasions", "Check availability"),
        ("MEET", "Spaces, capacities, equipment, catering", "Request a proposal"),
        ("EXPLORE", "Elbe, cycling, Magdeburg and local itineraries", "Plan the visit"),
    ]
    add_matrix(doc, ["JOURNEY", "WHAT THE GUEST NEEDS", "PRIMARY ACTION"], journey_rows, [1700, 4960, 2700], header_fill=BLUE, font_size=9.2)
    add_h2(doc, "Website deliverables")
    deliverables = [
        "Modern bilingual information architecture and visual design.",
        "Homepage with a persistent date-and-guest booking bar and visible direct-booking benefit.",
        "Room comparison plus a mobile-friendly restaurant menu, current hours and reservation flow.",
        "Living event calendar with individual event pages, price, time, menu, availability and automatic archiving.",
        "Separate short enquiry flows, with meeting capacity, layout, equipment and package information for decision-makers.",
        "Analytics, conversion events, consent setup, redirects, accessibility and performance QA.",
    ]
    for item in deliverables:
        add_bullet(doc, item, bullet_id)
    add_callout(doc, "Conversion principle", "Every important page should answer three questions: Is this right for me? Can I trust it? What do I do next?", fill=SAND, accent=GOLD)

    # PAGE 5
    add_page_break(doc)
    add_section_title(doc, "Google visibility", "Search, Maps, reviews and direct-booking presence", "The goal is stronger qualified visibility - not vanity traffic - for people looking to stay, dine, meet or celebrate in and around Magdeburg.")
    add_h2(doc, "1. Google Business Profile and Maps")
    for item in [
        "Audit ownership, primary and secondary categories, amenities, hours, restaurant information, services and booking links.",
        "Unify name, address, phone and website details across directories and key travel sources.",
        "Publish fresh hotel, room, food, terrace, event and meeting photography with descriptive captions.",
        "Create a monthly Google post plan for offers, seasonal dining, events and direct-booking benefits.",
        "Add tracked links so calls, website visits, directions and booking actions can be attributed.",
    ]:
        add_bullet(doc, item, bullet_id)
    add_h2(doc, "2. Review and rating growth")
    add_body(doc, "We will not manufacture or filter reviews. We will build an ethical guest-feedback system that makes it easy for real guests to leave honest feedback and gives the hotel a consistent response process.")
    for item in [
        "Post-stay and post-dining review requests by email, QR card or approved guest communication.",
        "Response templates for positive, neutral and negative feedback, adapted by a human.",
        "Monthly review themes: cleanliness, service, breakfast, location, restaurant and event experience.",
        "Escalation process for operational issues so reviews inform service improvements.",
    ]:
        add_bullet(doc, item, bullet_id)
    add_h2(doc, "3. Hotel presence on Google")
    add_body(doc, "We will audit whether the hotel's booking provider can connect rates and availability to Google Hotel Center. Where eligible, this can support free booking links, the official-site presentation and Hotel Ads. Price accuracy, feed quality and landing-page consistency remain essential.")

    # PAGE 6
    add_page_break(doc)
    add_section_title(doc, "SEO programme", "Build durable demand around real guest intent", "SEO will connect technical health, useful hospitality content and local authority. It is not a campaign of adding keywords to every paragraph.")
    seo_rows = [
        ("Technical SEO", "Crawl/index audit, redirects, canonical logic, sitemap, Core Web Vitals, mobile accessibility, image optimisation and broken links."),
        ("On-page SEO", "Clear titles and headings, internal linking, image descriptions, unique room/event copy and conversion-focused metadata."),
        ("Local SEO", "Magdeburg and Elbe-area relevance, consistent business data, local citations, maps signals and destination partnerships."),
        ("Structured data", "Hotel/local business, organisation, breadcrumb and event markup where supported and accurate."),
        ("Content clusters", "Hotel near Magdeburg, Elberadweg stays, riverside dining, celebrations, meetings, seasonal events and weekend itineraries."),
        ("Authority", "Useful local partnerships, PR opportunities, tourism links, supplier/event relationships and content worth referencing."),
    ]
    add_matrix(doc, ["WORKSTREAM", "WHAT WE WILL DO"], seo_rows, [2200, 7160], header_fill=NAVY, font_size=9.1)
    add_spacer(doc, 8)
    add_h2(doc, "Priority landing-page opportunities")
    for item in [
        "Hotel by the Elbe near Magdeburg",
        "Hotel for cyclists on the Elberadweg",
        "Riverside restaurant and seasonal dining",
        "Private celebrations and family events",
        "Meeting and seminar venue near Magdeburg",
        "Easter, Christmas, Valentine's and Mother's Day event pages",
    ]:
        add_bullet(doc, item, bullet_id)
    add_callout(doc, "SEO standard", "Create useful, accurate pages for guests first; optimise those pages so search engines can understand and surface them.", fill=PALE_BLUE, accent=BLUE)

    # PAGE 7
    add_page_break(doc)
    add_section_title(doc, "GEO and AI discovery", "Help answer engines understand and recommend ElbRivera", "GEO means Generative Engine Optimization. Local geographic targeting is handled through local SEO; GEO prepares the hotel's information for AI-assisted search and recommendation experiences.")
    add_h2(doc, "Our GEO approach")
    geo_steps = [
        ("Create an authoritative hotel entity.", "Keep the hotel's name, address, amenities, location, restaurant and direct-booking facts consistent across the website and trusted profiles."),
        ("Publish answer-ready information.", "Use concise room summaries, FAQs, event facts, policies, parking details, cycling information, accessibility details and meeting capacities."),
        ("Structure the information.", "Use valid schema, descriptive headings, indexable HTML and strong internal connections between rooms, dining, events and destination pages."),
        ("Demonstrate first-hand expertise.", "Publish original hotel photography, local recommendations, chef/event content and useful itineraries rather than generic destination text."),
        ("Build corroboration.", "Pursue accurate tourism, local business, event, cycling and hospitality mentions that support the same core facts."),
        ("Measure AI-assisted discovery.", "Track AI-referral traffic where identifiable, branded search growth, long-tail queries, cited landing pages and enquiry quality."),
    ]
    for lead, detail in geo_steps:
        add_numbered(doc, lead + " " + detail, number_id, bold_lead=lead)
    add_h2(doc, "What GEO does not mean")
    add_body(doc, "There is no special tag that guarantees inclusion in an AI Overview, AI Mode or another assistant. Google's current guidance says the same core SEO and people-first content practices remain relevant. GEO is therefore a disciplined extension of SEO: clearer entities, better evidence, structured facts and genuinely useful answers.")
    add_callout(doc, "Target outcome", "When a traveller asks for a quiet hotel near Magdeburg, a riverside restaurant, an Elberadweg stop or a meeting venue, ElbRivera should be easy for both search engines and AI systems to understand.", fill=SAND, accent=GOLD)

    # PAGE 8
    add_page_break(doc)
    add_section_title(doc, "Paid media", "Buy visibility where demand already exists", "Ads should accelerate high-intent demand and seasonal campaigns while the organic programme compounds over time.")
    ads_rows = [
        ("Google Search", "Hotel near Magdeburg, quiet hotel, Elbe hotel, restaurant and event/meeting searches", "Direct booking and enquiries"),
        ("Hotel Ads", "Date-specific hotel shoppers in Search and Maps, subject to Hotel Center/feed eligibility", "Booking-engine traffic"),
        ("Meta campaigns", "Restaurant events, celebrations, packages and remarketing audiences", "Awareness and reservations"),
        ("Brand protection", "Hotel-name and high-intent branded searches where OTAs or competitors compete", "More official-site clicks"),
        ("Remarketing", "Recent room, event and meeting-page visitors, with consent", "Return visits and assisted conversion"),
    ]
    add_matrix(doc, ["CHANNEL", "ROLE", "PRIMARY RESULT"], ads_rows, [1900, 4860, 2600], header_fill=BLUE, font_size=9.0)
    add_h2(doc, "Campaign management service")
    for item in [
        "Campaign and keyword architecture by revenue line, language, location and season.",
        "Dedicated landing pages, ad copy, creative adaptation and offer testing.",
        "Conversion tracking for booking-engine handoffs, calls, restaurant actions and enquiries.",
        "Negative-keyword control, geographic targeting, budget pacing and search-term review.",
        "Monthly reporting by commercial outcome, not only impressions and clicks.",
    ]:
        add_bullet(doc, item, bullet_id)
    add_callout(doc, "Budget principle", "Media spend is paid directly by the hotel and is separate from management fees. Budgets will be recommended only after demand, seasonality, booking economics and tracking readiness are reviewed.", fill=PALE_BLUE, accent=BLUE)

    # PAGE 9
    add_page_break(doc)
    add_section_title(doc, "Brand, content and social", "Make ElbRivera consistently visible between booking moments", "The website, Google presence and social channels should tell the same story and give guests a reason to return.")
    add_h2(doc, "Monthly content system")
    content_rows = [
        ("Rooms and direct booking", "Room stories, packages, quiet stays, parking and booking benefits"),
        ("Restaurant", "Seasonal dishes, chef stories, terrace, menus and non-resident dining"),
        ("Events", "Launch, reminder, last-availability and post-event content"),
        ("Celebrate and meet", "Real setups, capacities, menus, testimonials and planning guidance"),
        ("Place", "Elbe, cycling, Magdeburg itineraries and local collaborations"),
        ("Trust", "Guest reviews, team moments, FAQs and behind-the-scenes service"),
    ]
    add_matrix(doc, ["CONTENT PILLAR", "EXAMPLES"], content_rows, [2600, 6760], header_fill=NAVY, font_size=9.2)
    add_h2(doc, "Recommended operating cadence")
    for item in [
        "One monthly editorial and campaign calendar linked to hotel revenue priorities.",
        "A planned baseline of 8-12 feed posts, 4-8 short videos and event-led stories, adjusted to available assets and seasonality.",
        "Website and Google Business Profile updates from the same approved source information.",
        "Community and review-response guidelines with escalation for sensitive guest issues.",
        "Quarterly brand and offer review: what is resonating, what is converting and what needs repositioning.",
    ]:
        add_bullet(doc, item, bullet_id)
    add_body(doc, "Content production volume is finalised after asset access, shooting requirements, languages and approval responsibilities are confirmed.", size=9.5, color=MUTED, after=0)

    # PAGE 10
    add_page_break(doc)
    add_section_title(doc, "90-day launch plan", "Build the foundation, launch the system, then learn", "A staged rollout reduces risk and makes each decision traceable to a guest or business need.")
    roadmap_rows = [
        ("DAYS 1-15", "Discover and measure", "Stakeholder workshop; analytics/access audit; booking and enquiry baseline; guest journeys; content inventory; keyword and competitor research."),
        ("DAYS 16-35", "Position and design", "Messaging; information architecture; homepage/room/restaurant/event/meeting wireframes; visual direction; measurement plan."),
        ("DAYS 36-65", "Build and prepare growth", "Responsive development; CMS; booking integration; forms; schema; technical SEO; Google profile and review workflow; campaign setup."),
        ("DAYS 66-80", "Test and launch", "Content migration; proofreading; mobile/accessibility/performance QA; redirects; consent; analytics validation; staff training."),
        ("DAYS 81-90", "Optimise", "Early conversion review; search coverage; ads and landing-page tests; event calendar; first monthly content cycle; action report."),
    ]
    add_matrix(doc, ["TIMING", "PHASE", "KEY OUTPUTS"], roadmap_rows, [1500, 2260, 5600], header_fill=BLUE, font_size=8.9)
    add_h2(doc, "What we need from the hotel")
    for item in [
        "One decision-maker and clear response times for approvals.",
        "Access to the website, domain, booking provider, Google profiles, analytics and advertising accounts.",
        "Accurate room, restaurant, event, meeting, policy and pricing information.",
        "Permission to use approved hotel photography, reviews and brand assets.",
        "Operational commitment to fulfil offers, respond to enquiries and request reviews ethically.",
    ]:
        add_bullet(doc, item, bullet_id)

    # PAGE 11
    add_page_break(doc)
    add_section_title(doc, "Measurement", "Report what moves the business", "Success will be measured against an agreed baseline after account access. Early indicators show whether the system is working before revenue patterns fully mature.")
    kpi_rows = [
        ("Direct demand", "Booking-bar use, booking-engine handoffs, free/paid booking-link clicks, direct-booking completion where available"),
        ("Restaurant and events", "Reservation actions, event-detail engagement, phone/WhatsApp clicks and completed enquiries"),
        ("Meetings and celebrations", "Qualified leads, form completion, response time and lead-to-confirmed-event rate"),
        ("Search and Maps", "Qualified impressions, local actions, non-brand clicks, indexed landing pages and target-query coverage"),
        ("Reviews", "Request volume, review velocity, response rate, sentiment themes and operational issues"),
        ("Paid media", "Spend, qualified clicks, conversion rate, cost per lead/booking action and attributable revenue where available"),
        ("GEO", "AI-referral sessions where identifiable, cited/visited pages, branded search and long-tail discovery"),
        ("Experience", "Mobile performance, accessibility, broken journeys and form abandonment"),
    ]
    add_matrix(doc, ["BUSINESS AREA", "PRIMARY MEASURES"], kpi_rows, [2400, 6960], header_fill=NAVY, font_size=8.9)
    add_h2(doc, "Reporting rhythm")
    for item in [
        "Live or shared dashboard for agreed core metrics.",
        "Monthly performance summary with insights, actions and next-month priorities.",
        "Quarterly strategy review covering offers, channel mix, content and conversion opportunities.",
    ]:
        add_bullet(doc, item, bullet_id)

    # PAGE 12
    add_page_break(doc)
    add_section_title(doc, "Engagement options", "Choose the level of partnership", "Investment is scoped after the discovery and access audit. This avoids pricing an unknown technical environment or recommending media spend without conversion economics.")
    package_rows = [
        ("FOUNDATION", "Website redesign and build; analytics; technical/on-page SEO; core Google profile improvements; staff handover", "A modern platform ready to operate"),
        ("GROWTH - RECOMMENDED", "Everything in Foundation plus monthly SEO, GEO, Google/Maps, review operations, event content and paid-media management", "One connected acquisition system"),
        ("FULL PRESENCE", "Everything in Growth plus ongoing social strategy, production planning, publishing support, campaign creative and quarterly brand consulting", "An embedded digital growth partnership"),
    ]
    add_matrix(doc, ["OPTION", "INCLUDES", "BEST FOR"], package_rows, [2200, 4960, 2200], header_fill=BLUE, font_size=8.9)
    add_h2(doc, "Commercial framework")
    for item in [
        "One-time strategy, design and development fee based on approved scope.",
        "Monthly retainer based on selected growth and content workstreams.",
        "Advertising spend, photography, travel, third-party software and legal review are separate unless explicitly included.",
        "Initial term, ownership, access, support and cancellation conditions will be set out in the final statement of work.",
    ]:
        add_bullet(doc, item, bullet_id)
    add_callout(doc, "Recommended next step", "A 45-minute discovery session with the owner or commercial lead, followed by account/access review and a fixed statement of work.", fill=SAND, accent=GOLD)

    # PAGE 13
    add_page_break(doc)
    add_section_title(doc, "Why Veer + Inderpreet", "Two partners, one connected brand system", "We combine website delivery, content operations and brand consulting so the promise made in an ad is carried through to the website, booking journey and guest-facing content.")
    add_h2(doc, "Selected website work")
    website_links = [
        ("Hotel Metropolis", "https://hotelmetropolis.in/"),
        ("Daselb", "https://daselb.com/"),
        ("Ambur", "https://ambur.co.in/"),
        ("Synterra Technologies", "https://synterra-technologies.vercel.app/"),
        ("Crickroo", "https://www.crickroo.com/"),
        ("Hormone Nutrition Clinic", "https://www.hormonenutritionclinic.com/"),
        ("Tripund Technologies", "https://tripundtechnologies.in/"),
    ]
    for name, url in website_links:
        p = doc.add_paragraph()
        set_num(p, bullet_id)
        p.paragraph_format.space_after = Pt(3)
        add_hyperlink(p, name, url, color=BLUE)
    add_h2(doc, "Selected content-management work")
    for name, url in [
        ("Singh Laly Official", "https://www.instagram.com/singhlalyofficial/"),
        ("Intellia MIET", "https://www.instagram.com/intellia_miet/"),
    ]:
        p = doc.add_paragraph()
        set_num(p, bullet_id)
        p.paragraph_format.space_after = Pt(3)
        add_hyperlink(p, name, url, color=BLUE)
    add_h2(doc, "The close")
    add_callout(doc, "Our proposal", "Make ElbRivera easier to discover, easier to trust and easier to book - then keep the entire presence current, measurable and commercially focused.", fill=PALE_BLUE, accent=BLUE)
    add_spacer(doc, 8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run("Veer Pratap Singh + Inderpreet Singh"), size=12.5, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_hyperlink(p, "View our portfolio and discuss the proposal", "https://veer-pratap-singh-portfolio.vercel.app", color=GOLD)

    # PAGE 14 - references and assumptions
    add_page_break(doc)
    add_section_title(doc, "Proposal notes", "Assumptions, guardrails and sources", "This proposal is designed to support a transparent commercial discussion.")
    add_h2(doc, "Important assumptions")
    assumptions = [
        "The audit is based on publicly available pages and a live PageSpeed test; no internal analytics, booking revenue, occupancy or advertising data were accessed.",
        "Search, Maps, free booking links, ad auctions, review ratings and AI citations are controlled by platforms and users. Improvements can be influenced but not guaranteed.",
        "Review work will follow platform policies and will never include fabricated reviews, incentives tied to positive sentiment or suppression of legitimate negative feedback.",
        "GEO supports clarity and eligibility but does not guarantee AI inclusion. All final claims, prices, capacities, hours, menus, policies and translations require hotel approval.",
        "Privacy, cookie/consent and legal text should be reviewed against the hotel's operational and legal requirements.",
    ]
    for item in assumptions:
        add_bullet(doc, item, bullet_id, size=9.5, after=3)
    add_h2(doc, "Primary sources reviewed")
    sources = [
        ("Hotel ElbRivera homepage", "https://www.hotel-elbrivera.de/"),
        ("Hotel ElbRivera event calendar", "https://www.hotel-elbrivera.de/restaurant/veranstaltungen-feiern/veranstaltungskalender/"),
        ("Hotel ElbRivera enquiry page", "https://www.hotel-elbrivera.de/kontakt/"),
        ("Live PageSpeed report for the event calendar", "https://pagespeed.web.dev/analysis/https-www-hotel-elbrivera-de-restaurant-veranstaltungen-feiern-veranstaltungskalender/jhkro8sbmm?form_factor=mobile"),
        ("Google: Tips to improve local ranking", "https://support.google.com/business/answer/7091?hl=en"),
        ("Google: Hotel Ads and Hotel Center", "https://support.google.com/hotelprices/answer/11946932?hl=en"),
        ("Google: Free booking links", "https://support.google.com/hotelprices/answer/10472394?hl=en"),
        ("Google: SEO Starter Guide", "https://developers.google.com/search/docs/fundamentals/seo-starter-guide"),
        ("Google: AI features and your website", "https://developers.google.com/search/docs/appearance/ai-features"),
        ("Google: Event structured data", "https://developers.google.com/search/docs/appearance/structured-data/event"),
    ]
    for label, url in sources:
        p = doc.add_paragraph()
        set_num(p, bullet_id)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        add_hyperlink(p, label, url, color=BLUE, size=9.1)

    # Keep document fields updatable in Word/LibreOffice.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_document()
